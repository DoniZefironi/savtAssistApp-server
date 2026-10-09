"""Индексатор бота: разбор файлов (PDF, Word, Excel, картинки, старые форматы через
LibreOffice), нарезка на куски, запись эмбеддингов, переиндексация и очистка
осиротевших. Yandex (эмбеддинги, OCR, описание картинок) и LibreOffice подменены,
файлы создаются настоящими во временном каталоге."""
import io
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from sqlalchemy import select

from app.models.embedding import EMBEDDING_DIM, Embedding
from app.schemas.faq import FaqCategoryCreateIn, FaqEntryCreateIn
from app.schemas.kb import KbArticleCreateIn, KbCategoryCreateIn
from app.services import bot_indexer, yandex_service
from app.services.faq_service import FaqCategoryService, FaqEntryService
from app.services.kb_service import KbArticleService, KbCategoryService

VECTOR = [0.25] * EMBEDDING_DIM


@pytest.fixture(autouse=True)
def fakes(tmp_path, monkeypatch):
    root = tmp_path / "uploads"
    (root / "documents").mkdir(parents=True)
    monkeypatch.setattr(bot_indexer, "UPLOAD_ROOT", root)
    state = SimpleNamespace(embedded=[], ocr=[], described=[], fail_embedding_for=None, ocr_text="Текст со скана",
                            description="Схема подключения насоса")

    async def embed_document(text):
        if state.fail_embedding_for and state.fail_embedding_for in text:
            raise RuntimeError("Yandex недоступен")
        state.embedded.append(text)
        return VECTOR

    async def ocr_image(data):
        state.ocr.append(data)
        return state.ocr_text

    async def analyze_image(data, prompt, mime_type="image/png"):
        state.described.append(mime_type)
        return state.description

    async def no_wait(_seconds):
        return None

    monkeypatch.setattr(yandex_service, "embed_document", embed_document)
    monkeypatch.setattr(yandex_service, "ocr_image", ocr_image)
    monkeypatch.setattr(yandex_service, "analyze_image", analyze_image)
    monkeypatch.setattr(bot_indexer.asyncio, "sleep", no_wait)
    state.root = root
    return state


def _png(size=(300, 300)):
    buffer = io.BytesIO()
    Image.new("RGB", size, "white").save(buffer, format="PNG")
    return buffer.getvalue()


def _pdf_with_text(text):
    """Минимальный настоящий PDF с текстовым слоем."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


# --- нарезка ---

def test_chunks_overlap_and_empty_text():
    text = "а" * 2500

    chunks = bot_indexer._chunks(text)

    assert [len(c) for c in chunks] == [1200, 1200, 400]
    assert chunks[0][-150:] == chunks[1][:150]            # куски перекрываются
    assert bot_indexer._chunks("   \n ") == [] and bot_indexer._chunks("коротко") == ["коротко"]


# --- разбор файлов ---

async def test_pdf_text_layer_is_extracted(fakes):
    path = fakes.root / "documents" / "manual.pdf"
    path.write_bytes(_pdf_with_text("Pump control manual"))

    text = await bot_indexer._extract_text("/static/documents/manual.pdf")

    assert "Pump control manual" in text and fakes.ocr == [] and fakes.described == []


async def test_broken_pdf_gives_empty_text(fakes):
    (fakes.root / "documents" / "broken.pdf").write_bytes(b"not a pdf at all")

    assert await bot_indexer._extract_text("/static/documents/broken.pdf") == ""


def _fake_image(width=300, height=300, fmt="PNG"):
    return SimpleNamespace(data=b"img", image=SimpleNamespace(width=width, height=height, format=fmt))


class _FakePage:
    def __init__(self, text, images):
        self._text, self.images = text, images

    def extract_text(self):
        return self._text


async def test_pdf_images_are_ocred_on_scans_and_described_next_to_text(fakes, monkeypatch):
    pages = [
        _FakePage("", [_fake_image()]),                              # скан без текстового слоя — OCR
        _FakePage("Описание узла", [_fake_image(fmt="JPEG"), _fake_image(50, 50)]),   # схема рядом с текстом + иконка
    ]
    monkeypatch.setattr("pypdf.PdfReader", lambda path: SimpleNamespace(pages=pages))

    text = await bot_indexer._parse_pdf(Path("whatever.pdf"))

    assert "Текст со скана" in text and "Описание узла" in text
    assert "[Изображение на стр. 2]: Схема подключения насоса" in text
    assert len(fakes.ocr) == 1 and fakes.described == ["image/jpeg"]   # маленькая иконка пропущена


async def test_pdf_image_analysis_is_limited_and_failures_are_survived(fakes, monkeypatch):
    pages = [_FakePage("Текст", [_fake_image() for _ in range(20)])]
    monkeypatch.setattr("pypdf.PdfReader", lambda path: SimpleNamespace(pages=pages))

    await bot_indexer._parse_pdf(Path("big.pdf"))
    assert len(fakes.described) == bot_indexer._MAX_DOC_IMAGES_ANALYZED

    async def broken(*args, **kwargs):
        raise RuntimeError("модель недоступна")

    monkeypatch.setattr(yandex_service, "ocr_image", broken)
    monkeypatch.setattr(yandex_service, "analyze_image", broken)
    pages[:] = [_FakePage("", [_fake_image()]), _FakePage("Текст", [_fake_image()])]
    text = await bot_indexer._parse_pdf(Path("big.pdf"))
    assert "Текст" in text                                         # сбой картинки не теряет текст страницы


def test_image_helpers():
    assert bot_indexer._doc_image_mime_type(_fake_image(fmt="WEBP")) == "image/webp"
    assert bot_indexer._doc_image_mime_type(_fake_image(fmt="TIFF")) == "image/png"
    assert bot_indexer._doc_image_mime_type(SimpleNamespace(image=None)) == "image/png"
    assert bot_indexer._doc_image_too_small(_fake_image(100, 100)) is True
    assert bot_indexer._doc_image_too_small(_fake_image(100, 400)) is False
    assert bot_indexer._doc_image_too_small(SimpleNamespace(image=None)) is False
    assert bot_indexer._probe_image_bytes(_png((200, 120))) == (200, 120, "image/png")
    assert bot_indexer._probe_image_bytes(b"not an image") == (0, 0, "image/png")


def _make_docx(path, with_image=None):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Вводный абзац")
    table = doc.add_table(rows=2, cols=2)
    for row, values in enumerate((("Параметр", "Значение"), ("Давление", "6 бар"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    doc.add_paragraph("Заключение")
    if with_image:
        doc.add_picture(io.BytesIO(with_image))
    doc.save(str(path))


async def test_docx_keeps_the_order_of_text_and_tables(fakes):
    path = fakes.root / "documents" / "spec.docx"
    _make_docx(path)

    text = await bot_indexer._extract_text("/static/documents/spec.docx")

    assert text.splitlines() == ["Вводный абзац", "Параметр | Значение", "Давление | 6 бар", "Заключение"]


async def test_docx_pictures_are_described_unless_tiny(fakes):
    big, tiny = fakes.root / "documents" / "big.docx", fakes.root / "documents" / "tiny.docx"
    _make_docx(big, with_image=_png((300, 300)))
    _make_docx(tiny, with_image=_png((40, 40)))

    with_picture = await bot_indexer._parse_docx(big)
    without = await bot_indexer._parse_docx(tiny)

    assert "[Изображение в документе]: Схема подключения насоса" in with_picture
    assert "[Изображение" not in without
    assert fakes.described == ["image/png"]


async def test_unreadable_docx_picture_goes_through_libreoffice(fakes, monkeypatch):
    path = fakes.root / "documents" / "emf.docx"
    _make_docx(path)
    monkeypatch.setattr(bot_indexer, "_docx_image_blobs", lambda doc: [(b"EMF-bytes", "emf"), (b"WMF-bytes", "wmf")])
    converted = []

    async def convert(data, ext):
        converted.append(ext)
        return _png((300, 300)) if ext == "emf" else None

    monkeypatch.setattr(bot_indexer, "_convert_image_via_libreoffice", convert)

    text = await bot_indexer._parse_docx(path)

    assert converted == ["emf", "wmf"] and text.count("[Изображение в документе]") == 1


async def test_broken_docx_gives_empty_text(fakes):
    (fakes.root / "documents" / "broken.docx").write_bytes(b"PK-not-a-docx")

    assert await bot_indexer._parse_docx(fakes.root / "documents" / "broken.docx") == ""


async def test_excel_rows_become_text_lines(fakes):
    import openpyxl

    path = fakes.root / "documents" / "table.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Узел", "Мощность"])
    wb.active.append(["Насос", 7.5])
    wb.active.append([None, None])
    wb.save(str(path))

    assert bot_indexer._parse_excel(path) == "Узел | Мощность\nНасос | 7.5"
    assert bot_indexer._parse_excel(fakes.root / "documents" / "missing.xlsx") == ""


async def test_images_are_read_with_ocr_and_unknown_formats_are_skipped(fakes):
    (fakes.root / "documents" / "scan.jpg").write_bytes(_png())
    (fakes.root / "documents" / "archive.zip").write_bytes(b"zip")

    assert await bot_indexer._extract_text("/static/documents/scan.jpg") == "Текст со скана"
    assert await bot_indexer._extract_text("/static/documents/archive.zip") == ""
    assert await bot_indexer._extract_text("/static/documents/missing.pdf") == ""


async def test_legacy_office_files_are_converted_first(fakes, monkeypatch):
    (fakes.root / "documents" / "old.doc").write_bytes(b"binary doc")
    (fakes.root / "documents" / "dead.xls").write_bytes(b"binary xls")
    docx_path = fakes.root / "documents" / "converted.docx"
    _make_docx(docx_path)
    calls = []

    async def convert(path, target):
        calls.append((path.name, target))
        return docx_path.read_bytes() if path.name == "old.doc" else None

    monkeypatch.setattr(bot_indexer, "_convert_legacy_office", convert)

    text = await bot_indexer._extract_text("/static/documents/old.doc")
    failed = await bot_indexer._extract_text("/static/documents/dead.xls")

    assert "Вводный абзац" in text and failed == "" and calls == [("old.doc", "docx"), ("dead.xls", "xlsx")]


def test_libreoffice_runner_handles_every_outcome(tmp_path, monkeypatch):
    source = tmp_path / "report.doc"
    source.write_bytes(b"x")
    out_dir, profile = str(tmp_path / "out"), str(tmp_path / "profile")
    Path(out_dir).mkdir()

    monkeypatch.setattr(subprocess, "run", lambda *a, **kw: (_ for _ in ()).throw(subprocess.TimeoutExpired("soffice", 90)))
    assert bot_indexer._run_libreoffice_convert(source, "docx", out_dir, profile) is None

    monkeypatch.setattr(subprocess, "run", lambda *a, **kw: SimpleNamespace(returncode=1, stderr=b"boom"))
    assert bot_indexer._run_libreoffice_convert(source, "docx", out_dir, profile) is None

    monkeypatch.setattr(subprocess, "run", lambda *a, **kw: SimpleNamespace(returncode=0, stderr=b""))
    assert bot_indexer._run_libreoffice_convert(source, "docx", out_dir, profile) is None      # файла нет

    (Path(out_dir) / "report.docx").write_bytes(b"converted")
    assert bot_indexer._run_libreoffice_convert(source, "docx", out_dir, profile) == b"converted"


async def test_libreoffice_wrapper_survives_crashes(tmp_path, monkeypatch):
    source = tmp_path / "report.doc"
    source.write_bytes(b"x")

    def explode(*args, **kwargs):
        raise OSError("soffice не установлен")

    monkeypatch.setattr(bot_indexer, "_run_libreoffice_convert", explode)

    assert await bot_indexer._convert_legacy_office(source, "docx") is None


# --- запись эмбеддингов ---

async def _embeddings(db_session, source_type, source_id):
    return list((await db_session.execute(
        select(Embedding).where(Embedding.source_type == source_type, Embedding.source_id == source_id)
        .order_by(Embedding.chunk_index)
    )).scalars())


async def _faq(db_session, question="Как включить насос?", answer="Нажмите кнопку пуск"):
    category = await FaqCategoryService(db_session).create(FaqCategoryCreateIn(name=f"Кат {question}"))
    return await FaqEntryService(db_session).create(FaqEntryCreateIn(category_id=category.id, question=question, answer=answer))


async def test_faq_entry_is_indexed_and_reindexing_replaces_the_chunks(db_session, fakes):
    from app.models.faq_entry import FaqEntry
    entry = await db_session.get(FaqEntry, (await _faq(db_session)).id)

    await bot_indexer.index_faq_entry(db_session, entry)
    await bot_indexer.index_faq_entry(db_session, entry)

    rows = await _embeddings(db_session, "faq", entry.id)
    assert len(rows) == 1 and rows[0].meta == {"title": "Как включить насос?"}
    assert rows[0].content == "Вопрос: Как включить насос?\nОтвет: Нажмите кнопку пуск"


async def test_kb_article_text_and_readable_attachments_are_indexed(db_session, fakes):
    from app.models.kb_article_attachment import KbArticleAttachment
    from app.models.kbarticle import KbArticle
    category = await KbCategoryService(db_session).create(KbCategoryCreateIn(name="Общее"))
    created = await KbArticleService(db_session).create(
        KbArticleCreateIn(category_id=category.id, title="Запуск насоса", description="Шаг первый"))
    (fakes.root / "documents" / "scheme.jpg").write_bytes(_png())
    db_session.add_all([
        KbArticleAttachment(article_id=created.id, file_url="/static/documents/scheme.jpg", file_size_bytes=1,
                            doc_type="photo", mime_type="image/jpeg", title="Схема"),
        KbArticleAttachment(article_id=created.id, file_url="/static/documents/gone.pdf", file_size_bytes=1,
                            doc_type="pdf", mime_type="application/pdf", title="Нет файла"),
    ])
    await db_session.flush()

    await bot_indexer.index_kb_article(db_session, await db_session.get(KbArticle, created.id))

    [row] = await _embeddings(db_session, "kb_article", created.id)
    assert "Запуск насоса" in row.content and "Шаг первый" in row.content and "Текст со скана" in row.content


async def test_document_indexing_variants(db_session, fakes, make_document):
    ok = await make_document(file_url="/static/documents/ok.pdf", title="Паспорт")
    (fakes.root / "documents" / "ok.pdf").write_bytes(_pdf_with_text("Passport text"))
    empty = await make_document(file_url="/static/documents/missing.pdf", title="Без файла")
    internal = await make_document(file_url="/static/documents/ok.pdf", is_internal=True)
    db_session.add(Embedding(source_type="document", source_id=internal.id, content="старое", embedding=VECTOR, meta={}))
    await db_session.flush()

    await bot_indexer.index_document(db_session, ok)
    await bot_indexer.index_document(db_session, empty)
    await bot_indexer.index_document(db_session, internal)

    [good] = await _embeddings(db_session, "document", ok.id)
    assert "Passport text" in good.content and good.meta["extraction_failed"] is False
    [fallback] = await _embeddings(db_session, "document", empty.id)
    assert fallback.content == "Без файла" and fallback.meta["extraction_failed"] is True   # хотя бы заголовок
    assert await _embeddings(db_session, "document", internal.id) == []                   # служебный из поиска убран


# --- переиндексация ---

async def test_reindex_skips_what_is_done_unless_forced(db_session, fakes, make_document):
    await _faq(db_session)
    doc = await make_document(file_url="/static/documents/missing.pdf", title="Документ")

    first = await bot_indexer.reindex_all(db_session)
    second = await bot_indexer.reindex_all(db_session)
    forced = await bot_indexer.reindex_all(db_session, force=True)

    assert first["faq"] >= 1 and first["document"] >= 1
    assert second["faq"] == 0 and second["skipped"] >= 1
    assert second["document"] >= 1          # у документа не извлёкся текст — его пробуют снова
    assert forced["faq"] >= 1 and doc.id is not None


async def test_reindex_scope_and_project_limit(db_session, fakes, make_document, make_project, make_cabinet):
    await _faq(db_session)
    root, stranger = await make_project(), await make_project()
    cabinet = await make_cabinet(project_id=root.id)
    mine = await make_document(project_id=root.id, file_url="/static/documents/a.pdf")
    via_cabinet = await make_document(cabinet_id=cabinet.id, file_url="/static/documents/b.pdf")
    foreign = await make_document(project_id=stranger.id, file_url="/static/documents/c.pdf")

    only_faq = await bot_indexer.reindex_all(db_session, scope="faq")
    scoped = await bot_indexer.reindex_all(db_session, scope="document", project_id=root.id)

    assert only_faq["document"] == 0 and only_faq["faq"] >= 1
    assert scoped["document"] == 2 and scoped["faq"] == 0
    assert await _embeddings(db_session, "document", foreign.id) == []
    assert await _embeddings(db_session, "document", mine.id) and await _embeddings(db_session, "document", via_cabinet.id)


async def test_one_failure_does_not_stop_the_run(db_session, fakes):
    good = await _faq(db_session, "Рабочий вопрос", "ответ")
    await _faq(db_session, "СБОЙНЫЙ вопрос", "ответ")
    fakes.fail_embedding_for = "СБОЙНЫЙ"

    stats = await bot_indexer.reindex_all(db_session, scope="faq")

    assert stats["failed"] == 1 and stats["faq"] >= 1
    assert await _embeddings(db_session, "faq", good.id)


async def test_prune_removes_embeddings_of_vanished_sources(db_session, fakes):
    alive = await _faq(db_session)
    db_session.add_all([
        Embedding(source_type="faq", source_id=alive.id, content="жив", embedding=VECTOR, meta={}),
        Embedding(source_type="faq", source_id=987654, content="сирота", embedding=VECTOR, meta={}),
        Embedding(source_type="kb_article", source_id=987655, content="сирота", embedding=VECTOR, meta={}),
        Embedding(source_type="document", source_id=987656, content="сирота", embedding=VECTOR, meta={}),
        Embedding(source_type="other", source_id=1, content="неизвестный тип", embedding=VECTOR, meta={}),
    ])
    await db_session.flush()

    stats = await bot_indexer.prune_orphaned(db_session)

    assert stats == {"faq": 1, "kb_article": 1, "document": 1}
    assert await _embeddings(db_session, "faq", alive.id) and await _embeddings(db_session, "other", 1)


async def test_background_reindex_of_one_document(db_session, fakes, make_document, monkeypatch):
    class Ctx:
        async def __aenter__(self):
            return db_session

        async def __aexit__(self, *exc):
            return False

    pending = []
    monkeypatch.setattr("app.database.AsyncSessionLocal", lambda: Ctx())
    monkeypatch.setattr(bot_indexer, "spawn", pending.append)
    doc = await make_document(file_url="/static/documents/missing.pdf", title="Фоновый")

    bot_indexer.schedule_reindex_document(doc.id)
    bot_indexer.schedule_reindex_document(999999)       # нет такого документа — ничего не случается
    for task in pending:
        await task

    assert len(await _embeddings(db_session, "document", doc.id)) == 1
