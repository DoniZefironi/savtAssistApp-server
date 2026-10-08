"""Телеметрия ШУ: приём вебхука (история, текущее состояние, сигнал и пуш только
на реальное переключение названного бита), текущее состояние и история для
клиента и сотрудника, редактирование карты регистров, выгрузка в Excel,
автоочистка истории. Разбор битов и сборка карты — в test_telemetry_decoding.py.
Адреса регистров взяты с запасом (50000+), чтобы не пересекаться со
стандартной картой из миграций."""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import settings
from app.core.exceptions import AlreadyExistsError, NotFoundError, PermissionDeniedError
from app.models.audit_log import AuditLog
from app.models.cabinet_telemetry_event import CabinetTelemetryEvent
from app.models.notification import Notification
from app.schemas.telemetry import CabinetRegisterOverridePatchIn, RegisterDefinitionPatchIn
from app.services import notification_service, telemetry_service
from app.services.telemetry_service import (
    AdminRegisterMapService,
    TelemetryIngestService,
    UserTelemetryService,
    prune_old_telemetry_history,
    verify_telemetry_secret,
)

ADDR = 50000


@pytest.fixture
def env(monkeypatch):
    e = SimpleNamespace(signals=[], pushes=[])

    async def signal(cabinet_id, event_id):
        e.signals.append((cabinet_id, event_id))

    async def push(session, user_id, title, body, data=None, notification_type=None):
        e.pushes.append((user_id, notification_type))

    monkeypatch.setattr(telemetry_service, "publish_telemetry_event", signal)
    monkeypatch.setattr(notification_service, "send_push", push)
    return e


async def _admin_user(make_user):
    return (await make_user("admin")).id


async def _name_bits(db_session, make_user, cabinet):
    actor = await _admin_user(make_user)
    svc = AdminRegisterMapService(db_session)
    await svc.create_override(cabinet.id, ADDR, 0, "Авария насоса", None, actor, "admin")
    await svc.create_override(cabinet.id, ADDR, 3, "Перегрев", "Датчик T1", actor, "admin")


# --- секрет вебхука ---

def test_webhook_secret(monkeypatch):
    monkeypatch.setattr(settings, "telemetry_webhook_secret", "s3cret")
    assert verify_telemetry_secret("s3cret") is True
    assert verify_telemetry_secret("wrong") is False and verify_telemetry_secret(None) is False
    monkeypatch.setattr(settings, "telemetry_webhook_secret", "")
    assert verify_telemetry_secret("") is False  # пустой секрет не открывает вебхук


# --- цели подключения ---

async def test_targets_list_only_complete_live_cabinets(db_session, env, make_cabinet):
    complete = await make_cabinet(mqtt_topic="a/1", mqtt_host="h", mqtt_port=1883, mqtt_username="u", mqtt_password="p")
    await make_cabinet(mqtt_topic="a/2", mqtt_host="h")  # без порта
    await make_cabinet(mqtt_topic="a/3", mqtt_host="h", mqtt_port=1883, deleted_at=datetime.now(timezone.utc))

    targets = await TelemetryIngestService(db_session).list_targets()

    ours = [t for t in targets if t.cabinet_id == complete.id]
    assert len(ours) == 1 and (ours[0].host, ours[0].port, ours[0].topic) == ("h", 1883, "a/1")
    assert ours[0].password == "p"
    assert all(t.cabinet_id != 0 and t.topic for t in targets)


# --- приём ---

async def test_ingest_stores_history_and_state(db_session, env, make_cabinet):
    cabinet = await make_cabinet(mqtt_topic="t/ingest")
    svc = TelemetryIngestService(db_session)
    stamp = datetime(2026, 9, 1, 12, 0, tzinfo=timezone.utc)

    await svc.ingest("t/ingest", {ADDR: 5, ADDR + 1: 0}, stamp)
    await svc.ingest("t/ingest", {ADDR: 7}, None)

    events = (await db_session.execute(
        select(CabinetTelemetryEvent).where(CabinetTelemetryEvent.cabinet_id == cabinet.id).order_by(CabinetTelemetryEvent.id)
    )).scalars().all()
    assert len(events) == 2 and events[0].received_at == stamp
    assert events[0].raw_payload == {str(ADDR): 5, str(ADDR + 1): 0}
    state = {s.address: s.value for s in await svc.state_repo.list_for_cabinet(cabinet.id)}
    assert state == {ADDR: 7, ADDR + 1: 0}  # последнее значение каждого адреса


async def test_ingest_unknown_topic(db_session, env):
    with pytest.raises(NotFoundError):
        await TelemetryIngestService(db_session).ingest("no/such", {ADDR: 1}, None)


async def test_signal_and_push_only_on_named_bit_transition(
    db_session, env, make_user, make_cabinet, make_project, link_user_project,
):
    project = await make_project()
    cabinet = await make_cabinet(mqtt_topic="t/alarm", project_id=project.id, admin_internal_name="Насосная")
    member, outsider = await make_user(), await make_user()
    await link_user_project(member, project)
    await _name_bits(db_session, make_user, cabinet)
    svc = TelemetryIngestService(db_session)

    await svc.ingest("t/alarm", {ADDR: 0b0000}, None)   # первое сообщение, всё в нуле
    assert env.signals == [] and env.pushes == []
    await svc.ingest("t/alarm", {ADDR: 0b0010}, None)   # переключился безымянный бит 1
    assert env.signals == [] and env.pushes == []
    await svc.ingest("t/alarm", {ADDR: 0b0011}, None)   # взведён названный бит 0 — авария
    assert len(env.signals) == 1 and env.pushes == [(member.id, "cabinet_alarm")]
    await svc.ingest("t/alarm", {ADDR: 0b0011}, None)   # повтор того же состояния
    assert len(env.signals) == 1 and len(env.pushes) == 1

    [note] = (await db_session.execute(select(Notification).where(Notification.user_id == member.id))).scalars().all()
    assert note.title == "Авария ШУ «Насосная»" and note.body == "Авария насоса"
    assert note.data == {"cabinet_id": cabinet.id}
    assert (await db_session.execute(select(Notification).where(Notification.user_id == outsider.id))).first() is None


async def test_clearing_an_alarm_signals_but_does_not_push(db_session, env, make_user, make_cabinet, make_project, link_user_project):
    project = await make_project()
    cabinet = await make_cabinet(mqtt_topic="t/clear", project_id=project.id)
    member = await make_user()
    await link_user_project(member, project)
    await _name_bits(db_session, make_user, cabinet)
    svc = TelemetryIngestService(db_session)
    await svc.ingest("t/clear", {ADDR: 0b0001}, None)
    env.signals.clear()
    env.pushes.clear()

    await svc.ingest("t/clear", {ADDR: 0b0000}, None)

    assert len(env.signals) == 1 and env.pushes == []


async def test_each_new_alarm_bit_is_a_separate_push(db_session, env, make_user, make_cabinet, make_project, link_user_project):
    project = await make_project()
    cabinet = await make_cabinet(mqtt_topic="t/two", project_id=project.id)
    member = await make_user()
    await link_user_project(member, project)
    await _name_bits(db_session, make_user, cabinet)

    await TelemetryIngestService(db_session).ingest("t/two", {ADDR: 0b1001}, None)

    bodies = sorted(n.body for n in (await db_session.execute(select(Notification).where(Notification.user_id == member.id))).scalars())
    assert bodies == ["Авария насоса", "Перегрев"]


async def test_alarm_without_members_is_silent(db_session, env, make_user, make_cabinet):
    cabinet = await make_cabinet(mqtt_topic="t/nobody")
    await _name_bits(db_session, make_user, cabinet)

    await TelemetryIngestService(db_session).ingest("t/nobody", {ADDR: 0b0001}, None)

    assert len(env.signals) == 1 and env.pushes == []


# --- текущее состояние и история ---

async def test_current_state_for_member_and_stranger(
    db_session, env, make_user, make_cabinet, make_project, link_user_project,
):
    project = await make_project()
    cabinet = await make_cabinet(mqtt_topic="t/state", project_id=project.id)
    member, stranger = await make_user(), await make_user()
    await link_user_project(member, project)
    await _name_bits(db_session, make_user, cabinet)
    await TelemetryIngestService(db_session).ingest("t/state", {ADDR: 0b1001}, None)
    svc = UserTelemetryService(db_session)

    state = await svc.get_current_state(member.id, cabinet.id)
    full = await svc.get_current_state(member.id, cabinet.id, include_unnamed=True)
    admin_view = await svc.get_current_state_admin(cabinet.id)

    assert sorted((r.bit, r.name, r.value) for r in state.registers) == [(0, "Авария насоса", 1), (3, "Перегрев", 1)]
    assert len(full.registers) == 16
    assert len(admin_view.registers) == 2
    with pytest.raises(PermissionDeniedError):
        await svc.get_current_state(stranger.id, cabinet.id)
    with pytest.raises(NotFoundError):
        await svc.get_current_state_admin(999999)


async def test_history_hides_messages_without_named_bits(db_session, env, make_user, make_cabinet):
    cabinet = await make_cabinet(mqtt_topic="t/history")
    await _name_bits(db_session, make_user, cabinet)
    svc = TelemetryIngestService(db_session)
    await svc.ingest("t/history", {ADDR: 0b0100}, None)   # только безымянный бит
    await svc.ingest("t/history", {ADDR: 0b0001}, None)   # названный
    reader = UserTelemetryService(db_session)

    page = await reader.list_history_for_cabinet_admin(cabinet.id, 1, 10)
    full = await reader.list_history_for_cabinet_admin(cabinet.id, 1, 10, include_unnamed=True)

    assert page.total == 2 and len(page.items) == 1  # пагинация по сырым событиям
    assert len(full.items) == 2
    with pytest.raises(NotFoundError):
        await reader.list_history_for_cabinet_admin(999999, 1, 10)


async def test_prune_old_history_keeps_state(db_session, env, make_cabinet):
    cabinet = await make_cabinet(mqtt_topic="t/prune")
    svc = TelemetryIngestService(db_session)
    old = datetime.now(timezone.utc) - timedelta(days=30)
    await svc.ingest("t/prune", {ADDR: 1}, old)
    await svc.ingest("t/prune", {ADDR: 2}, None)

    deleted = await prune_old_telemetry_history(db_session, retention_days=14)

    left = (await db_session.execute(select(CabinetTelemetryEvent).where(CabinetTelemetryEvent.cabinet_id == cabinet.id))).scalars().all()
    assert deleted >= 1 and len(left) == 1
    assert [s.value for s in await svc.state_repo.list_for_cabinet(cabinet.id)] == [2]


async def test_prune_uses_configured_retention_by_default(db_session, env, make_cabinet, monkeypatch):
    cabinet = await make_cabinet(mqtt_topic="t/prune2")
    svc = TelemetryIngestService(db_session)
    await svc.ingest("t/prune2", {ADDR: 1}, datetime.now(timezone.utc) - timedelta(days=3))
    monkeypatch.setattr(settings, "telemetry_history_retention_days", 2)

    await prune_old_telemetry_history(db_session)

    assert (await db_session.execute(select(CabinetTelemetryEvent).where(CabinetTelemetryEvent.cabinet_id == cabinet.id))).first() is None


async def test_cabinet_telemetry_access_check(db_session, env, make_user, make_cabinet, link_user_cabinet, monkeypatch):
    class Ctx:
        async def __aenter__(self):
            return db_session

        async def __aexit__(self, *exc):
            return False

    monkeypatch.setattr("app.database.AsyncSessionLocal", lambda: Ctx())
    owner, stranger = await make_user(), await make_user()
    cabinet = await make_cabinet()
    await link_user_cabinet(owner, cabinet)

    assert await telemetry_service.check_cabinet_telemetry_access(cabinet.id, owner.id)
    assert not await telemetry_service.check_cabinet_telemetry_access(cabinet.id, stranger.id)


# --- карта регистров ---

async def test_standard_map_crud(db_session, env, make_user):
    actor = await _admin_user(make_user)
    svc = AdminRegisterMapService(db_session)

    created = await svc.create_definition(ADDR + 10, 2, "Нет воды", "Датчик сухого хода", actor, "admin")
    with pytest.raises(AlreadyExistsError):
        await svc.create_definition(ADDR + 10, 2, "Дубль", None, actor, "admin")
    other = await svc.create_definition(ADDR + 10, 4, "Нет фазы", None, actor, "admin")

    renamed = await svc.update_definition(created.id, RegisterDefinitionPatchIn(name="Сухой ход"), actor, "admin")
    assert renamed.name == "Сухой ход" and renamed.description == "Датчик сухого хода"
    with pytest.raises(AlreadyExistsError):
        await svc.update_definition(other.id, RegisterDefinitionPatchIn(bit=2), actor, "admin")  # занято
    moved = await svc.update_definition(other.id, RegisterDefinitionPatchIn(bit=5), actor, "admin")
    assert moved.bit == 5
    same = await svc.update_definition(other.id, RegisterDefinitionPatchIn(bit=5, name="Нет фазы L1"), actor, "admin")
    assert same.name == "Нет фазы L1"  # собственное место занятым не считается

    assert {ADDR + 10} == {d.address for d in await svc.list_definitions() if d.address == ADDR + 10}
    await svc.delete_definition(created.id, actor, "admin")
    for call in (lambda: svc.update_definition(created.id, RegisterDefinitionPatchIn(name="x"), actor, "admin"),
                 lambda: svc.delete_definition(created.id, actor, "admin")):
        with pytest.raises(NotFoundError):
            await call()
    actions = [a.action for a in (await db_session.execute(select(AuditLog).where(AuditLog.action.like("register_definition.%")))).scalars()]
    assert sorted(set(actions)) == ["register_definition.create", "register_definition.delete", "register_definition.update"]


async def test_override_crud_is_scoped_to_its_cabinet(db_session, env, make_user, make_cabinet):
    actor = await _admin_user(make_user)
    mine, other = await make_cabinet(), await make_cabinet()
    svc = AdminRegisterMapService(db_session)

    created = await svc.create_override(mine.id, ADDR, 1, "Своя авария", None, actor, "admin")
    with pytest.raises(AlreadyExistsError):
        await svc.create_override(mine.id, ADDR, 1, "Дубль", None, actor, "admin")
    await svc.create_override(other.id, ADDR, 1, "Та же в другом шкафу", None, actor, "admin")  # другому ШУ можно
    second = await svc.create_override(mine.id, ADDR, 2, "Вторая", None, actor, "admin")

    assert [o.name for o in await svc.list_overrides(mine.id)] == ["Своя авария", "Вторая"]
    updated = await svc.update_override(mine.id, created.id, CabinetRegisterOverridePatchIn(name="Новая"), actor, "admin")
    assert updated.name == "Новая"
    with pytest.raises(AlreadyExistsError):
        await svc.update_override(mine.id, second.id, CabinetRegisterOverridePatchIn(bit=1), actor, "admin")
    with pytest.raises(NotFoundError):
        await svc.update_override(other.id, created.id, CabinetRegisterOverridePatchIn(name="x"), actor, "admin")  # чужой шкаф
    with pytest.raises(NotFoundError):
        await svc.delete_override(other.id, created.id, actor, "admin")
    await svc.delete_override(mine.id, created.id, actor, "admin")
    assert [o.name for o in await svc.list_overrides(mine.id)] == ["Вторая"]
    for call in (lambda: svc.list_overrides(999999),
                 lambda: svc.create_override(999999, ADDR, 0, "x", None, actor, "admin")):
        with pytest.raises(NotFoundError):
            await call()


# --- выгрузка в Excel ---

def _rows(xlsx: bytes):
    ws = load_workbook(BytesIO(xlsx)).active
    return ws.title, [[c.value for c in row] for row in ws.iter_rows()]


async def test_standard_map_export(db_session, env, make_user):
    actor = await _admin_user(make_user)
    svc = AdminRegisterMapService(db_session)
    await svc.create_definition(ADDR + 20, 7, "Бит семь", "Описание", actor, "admin")
    await svc.create_definition(ADDR + 20, 1, "Бит один", None, actor, "admin")

    title, rows = _rows(await svc.export_definitions_xlsx())

    assert title == "Карта регистров" and rows[0] == ["Адрес", "Бит", "Название", "Описание"]
    ours = [r for r in rows if r[0] == ADDR + 20]
    assert ours == [[ADDR + 20, 1, "Бит один", None], [ADDR + 20, 7, "Бит семь", "Описание"]]  # по адресу и биту


async def test_cabinet_map_export_marks_source(db_session, env, make_user, make_cabinet):
    actor = await _admin_user(make_user)
    cabinet = await make_cabinet(admin_internal_name="Насосная")
    svc = AdminRegisterMapService(db_session)
    await svc.create_definition(ADDR + 30, 0, "Стандартное имя", None, actor, "admin")
    await svc.create_definition(ADDR + 30, 1, "Только стандарт", None, actor, "admin")
    await svc.create_override(cabinet.id, ADDR + 30, 0, "Имя для этого ШУ", "своё", actor, "admin")

    title, rows = _rows(await svc.export_cabinet_map_xlsx(cabinet.id))

    assert title == "Карта регистров Насосная" and rows[0][-1] == "Источник"
    ours = {(r[0], r[1]): r for r in rows[1:] if r[0] == ADDR + 30}
    assert ours[(ADDR + 30, 0)][2:] == ["Имя для этого ШУ", "своё", "Переопределено для этого ШУ"]
    assert ours[(ADDR + 30, 1)][2:] == ["Только стандарт", None, "Стандартная карта"]
    with pytest.raises(NotFoundError):
        await svc.export_cabinet_map_xlsx(999999)


async def test_export_sheet_title_is_cut_to_excel_limit(db_session, env, make_cabinet):
    cabinet = await make_cabinet(admin_internal_name="Очень длинное название шкафа управления насосной станцией")

    title, _ = _rows(await AdminRegisterMapService(db_session).export_cabinet_map_xlsx(cabinet.id))

    assert len(title) == 31
